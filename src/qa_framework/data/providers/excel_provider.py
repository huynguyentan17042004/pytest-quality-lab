from pathlib import Path

from openpyxl import load_workbook

from qa_framework.exceptions import DataValidationError


class ExcelDataProvider:
    def load(self, path: Path, sheet: str | None = None, **_kwargs) -> list[dict]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet and sheet not in workbook.sheetnames:
                raise DataValidationError(f"{path}: sheet {sheet!r} not found")
            worksheet = workbook[sheet] if sheet else workbook.active
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers or not all(headers):
                raise DataValidationError(f"{path}/{worksheet.title}: header row is required")
            return [
                dict(zip(headers, row, strict=False))
                for row in rows
                if any(value is not None for value in row)
            ]
        finally:
            workbook.close()
